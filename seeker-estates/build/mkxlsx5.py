import sys; sys.path.insert(0,'.')
from properties import PROPS, VENUES, FACTS
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
F="Arial"; NAVY="1F3A5F"
thin=Side(style="thin",color="CCCCCC"); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
path="/home/claude/seeker-estates/outputs/Seeker-Estates-Content-Engine.xlsx"
wb=load_workbook(path)
if "Properties" in wb.sheetnames: del wb["Properties"]
ws=wb.create_sheet("Properties",1)
def header(row,cols,widths):
    for i,(c,w) in enumerate(zip(cols,widths),1):
        cell=ws.cell(row=row,column=i,value=c); cell.font=Font(name=F,bold=True,color="FFFFFF",size=10)
        cell.fill=PatternFill("solid",fgColor=NAVY); cell.alignment=Alignment(vertical="center",horizontal="center",wrap_text=True)
        cell.border=BOX; ws.column_dimensions[get_column_letter(i)].width=w
    ws.row_dimensions[row].height=30
ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=11)
c=ws.cell(row=1,column=1,value="Properties — from Corentin's pack, 10 Sep 2026"); c.font=Font(name=F,bold=True,size=15,color=NAVY)
ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=11)
c=ws.cell(row=2,column=1,value="19 properties. Rates are Corentin's low / high / base figures. Journey times in the venue map are estimates — verify on TfL before any post states one."); c.font=Font(name=F,size=9,italic=True,color="666666")
cols=["Ref","Postcode","Area","Station","Type","Sleeps","Low £","High £","Base £","Nearest venue(s)","Notes"]
header(4,cols,[6,10,17,26,7,7,7,7,7,32,44])
r=5
for p in PROPS:
    for i,v in enumerate(p,1):
        c=ws.cell(row=r,column=i,value=v); c.font=Font(name=F,size=9); c.border=BOX
        c.alignment=Alignment(vertical="top",wrap_text=(i in (4,10,11)))
        if i in (7,8,9): c.number_format="#,##0"
    ws.row_dimensions[r].height=30; r+=1
last=r-1
ws.cell(row=r,column=5,value="Avg base").font=Font(name=F,size=9,bold=True)
c=ws.cell(row=r,column=9,value=f"=AVERAGE(I5:I{last})"); c.number_format="#,##0"; c.font=Font(name=F,size=9,bold=True)
ws.cell(row=r,column=5).border=BOX; c.border=BOX
ws.cell(row=r,column=6,value=f"=SUM(F5:F{last})").font=Font(name=F,size=9,bold=True); ws.cell(row=r,column=6).border=BOX
ws.cell(row=r,column=4,value="Total sleeps →").font=Font(name=F,size=9,italic=True,color="666666")
r+=3
ws.cell(row=r,column=1,value="VENUE MAP — which properties to feature for which event").font=Font(name=F,bold=True,size=12,color=NAVY); r+=1
header(r,["Venue","Best-placed properties (verify journey times)"],[6,10]); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=11); r+=1
for v,ps in VENUES.items():
    ws.cell(row=r,column=1,value=v).font=Font(name=F,size=9,bold=True)
    ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=11)
    c=ws.cell(row=r,column=2,value=" · ".join(ps)); c.font=Font(name=F,size=9); c.alignment=Alignment(wrap_text=True,vertical="top")
    for i in range(1,12): ws.cell(row=r,column=i).border=BOX
    ws.row_dimensions[r].height=28; r+=1
r+=2
ws.cell(row=r,column=1,value="BUSINESS FACTS the team writes from").font=Font(name=F,bold=True,size=12,color=NAVY); r+=1
for k,v in FACTS.items():
    ws.cell(row=r,column=1,value=k.replace("_"," ")).font=Font(name=F,size=9,bold=True)
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=3)
    ws.merge_cells(start_row=r,start_column=4,end_row=r,end_column=11)
    c=ws.cell(row=r,column=4,value=v); c.font=Font(name=F,size=9); c.alignment=Alignment(wrap_text=True)
    for i in range(1,12): ws.cell(row=r,column=i).border=BOX
    r+=1
ws.freeze_panes="A5"; ws.sheet_view.showGridLines=False
wb.save(path); print("properties tab added")
