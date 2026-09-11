import sys; sys.path.insert(0,'/home/claude/seeker-estates/build')
from datetime import date, timedelta
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
F="Arial"; NAVY="1F3A5F"
thin=Side(style="thin",color="CCCCCC"); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
wb=load_workbook("/home/claude/seeker-estates/build/_c3.xlsx")
def header(ws,row,cols,widths):
    for i,(c,w) in enumerate(zip(cols,widths),1):
        cell=ws.cell(row=row,column=i,value=c); cell.font=Font(name=F,bold=True,color="FFFFFF",size=10)
        cell.fill=PatternFill("solid",fgColor=NAVY); cell.alignment=Alignment(vertical="center",horizontal="center",wrap_text=True)
        cell.border=BOX; ws.column_dimensions[get_column_letter(i)].width=w
    ws.row_dimensions[row].height=34
def title(ws,t,s,n):
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=n); c=ws.cell(row=1,column=1,value=t)
    c.font=Font(name=F,bold=True,size=15,color=NAVY); ws.row_dimensions[1].height=26
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=n); c=ws.cell(row=2,column=1,value=s)
    c.font=Font(name=F,size=9,italic=True,color="666666"); ws.row_dimensions[2].height=16

# ---------- SHEET 4: WEEKLY TRACKER ----------
ws=wb.create_sheet("Weekly Tracker")
cols=["Week","W/C","Posts published","Total reach","% reach from non-followers","New followers","Saves","Shares","Profile visits","Link clicks","Guest enquiries","Bookings from social","Landlord enquiries","Landlord calls booked","Notes / what worked"]
widths=[7,11,10,12,13,11,9,9,11,10,12,12,12,12,46]
title(ws,"Weekly Tracker","Fill the yellow cells every Monday for the week just gone. Everything else calculates. Track landlord enquiries SEPARATELY — never blend them into guest numbers.",len(cols))
header(ws,4,cols,widths)
start=date(2026,9,14); r=5
for w in range(1,14):
    wc=start+timedelta(weeks=w-1)
    ws.cell(row=r,column=1,value=f"W{w}").font=Font(name=F,size=9,bold=True)
    c=ws.cell(row=r,column=2,value=wc); c.number_format="ddd dd mmm"; c.font=Font(name=F,size=9)
    for i in range(3,15):
        c=ws.cell(row=r,column=i); c.fill=PatternFill("solid",fgColor="FFFFCC"); c.font=Font(name=F,size=9)
        c.border=BOX
        if i==5: c.number_format="0%"
    for i in (1,2,15):
        ws.cell(row=r,column=i).border=BOX
    ws.cell(row=r,column=15).alignment=Alignment(vertical="top",wrap_text=True)
    ws.cell(row=r,column=15).font=Font(name=F,size=9)
    ws.row_dimensions[r].height=30
    r+=1
lastr=r-1
# totals row
ws.cell(row=r,column=1,value="TOTAL").font=Font(name=F,bold=True,size=10,color="FFFFFF")
ws.cell(row=r,column=1).fill=PatternFill("solid",fgColor=NAVY)
for i in list(range(3,5))+list(range(6,15)):
    c=ws.cell(row=r,column=i,value=f'=SUM({get_column_letter(i)}5:{get_column_letter(i)}{lastr})')
    c.font=Font(name=F,bold=True,size=9); c.border=BOX; c.fill=PatternFill("solid",fgColor="E8EEF4")
c=ws.cell(row=r,column=5,value=f'=IFERROR(AVERAGE(E5:E{lastr}),0)'); c.number_format="0%"
c.font=Font(name=F,bold=True,size=9); c.border=BOX; c.fill=PatternFill("solid",fgColor="E8EEF4")
tot=r

# derived metrics block
r=tot+3
ws.cell(row=r,column=1,value="DERIVED — the numbers that actually matter").font=Font(name=F,bold=True,size=13,color=NAVY)
r+=1
header(ws,r,["Metric","Value","What good looks like at 90 days","Why"],[34,14,40,60])
M=[
("Avg reach per post",f'=IFERROR(D{tot}/C{tot},0)',"3,000–8,000 for a new local account","Reach per post is the only honest measure of whether the content is working. Follower count lags it by months."),
("Save rate (saves ÷ reach)",f'=IFERROR(G{tot}/D{tot},0)',"Above 1.5% is strong","Saves are the strongest quality signal on both IG and TikTok. Chase this, not likes."),
("Share rate (shares ÷ reach)",f'=IFERROR(H{tot}/D{tot},0)',"Above 0.8% is strong","A share is someone putting their own reputation behind your post. It's the highest-value action there is."),
("Profile visit rate",f'=IFERROR(I{tot}/D{tot},0)',"Above 2%","Measures whether the content makes people curious about who made it. Low = content is good but brand is invisible."),
("Link click-through",f'=IFERROR(J{tot}/I{tot},0)',"Above 8% of profile visits","Tests the bio and the landing page, not the content."),
("Guest enquiry → booking",f'=IFERROR(L{tot}/K{tot},0)',"Above 20%","Tests your response speed and your rates, not your marketing."),
("Landlord enquiry → call",f'=IFERROR(N{tot}/M{tot},0)',"Above 40%","Landlord leads are warm by the time they DM. If this is low, the follow-up is broken."),
("Cost per landlord lead","", "Under £150 organic","One landlord conversion is worth more than a year of bookings. Value these leads properly."),
]
r+=1
for m in M:
    for i,v in enumerate(m,1):
        c=ws.cell(row=r,column=i,value=v); c.font=Font(name=F,size=9); c.border=BOX
        c.alignment=Alignment(vertical="top",wrap_text=(i in (3,4)))
        if i==1: c.font=Font(name=F,size=9,bold=True)
        if i==2:
            c.fill=PatternFill("solid",fgColor="E8EEF4")
            c.number_format=("0.00%" if "rate" in m[0].lower() or "→" in m[0] else "#,##0")
    ws.row_dimensions[r].height=40; r+=1

# ---------- SHEET 5: EVENT MAP ----------
ws=wb.create_sheet("Event Map")
cols=["Date(s)","Event","Venue / area","Nearest tube/rail","Audience","Demand","Content angle already in the calendar"]
widths=[22,42,28,24,26,11,52]
title(ws,"Event map — the 26 dates that drive demand","Sourced from the full research file. Red = peak demand week, blue = trough. Use this for pricing as much as for content.",len(cols))
header(ws,4,cols,widths)
E=[
("10 Sep 2026 – 11 Jul 2027","Bayeux Tapestry","British Museum, WC1","Holborn / Russell Sq","International, family, culture","HIGH","Launch carousel 11 Sep — a 10-month standing reason to visit London"),
("12–20 Sep 2026","London Design Festival","11 districts, London-wide","Various","International, corporate","HIGH","Free-things reel 12 Sep — dispersed across zones, favours non-Zone-1 stock"),
("17–21 Sep 2026","London Fashion Week SS27","W1 / Strand / Soho","Oxford Circus","International, corporate","PEAK","Weekly anchor 14 Sep. Buyers book 5–7 nights"),
("25–27 Sep 2026","Laver Cup","The O2, SE10","North Greenwich","International, corporate","PEAK","Anchor 21 Sep — plus the 'ATP Finals are in Turin' correction"),
("26 Sep 2026","England v Spain","Wembley Stadium","Wembley Park","International, domestic","HIGH","Matchday food guide 26 Sep"),
("1–11 Oct 2026","BFI London Film Festival","Southbank + citywide","Waterloo","International, corporate","HIGH","Screenings carousel 3 Oct. Delegates take 7–11 nights"),
("4 & 11 Oct 2026","NFL London (Tottenham)","Tottenham Hotspur Stadium","Seven Sisters","International (US)","HIGH","Anchor 5 Oct + 'six things that will confuse you' 6 Oct"),
("14–18 Oct 2026","FRIEZE LONDON + MASTERS","The Regent's Park, NW1","Regent's Park","International HNW","PEAK ★","Anchor 12 Oct + free-Frieze counter-programming 13 Oct"),
("18 Oct 2026","NFL London (Wembley)","Wembley Stadium","Wembley Park","International (US)","PEAK","Inside the Frieze-week anchor. US visitors book 4–7 nights"),
("19–23 Oct 2026","— POST-FRIEZE AIR POCKET —","London-wide","—","—","TROUGH","Value week: anchor 19 Oct + honest price comparison 20 Oct"),
("26–30 Oct 2026","Autumn half term","London-wide","—","Family, domestic","HIGH","Free-things-with-kids reel posted Sat 24 Oct, before it starts"),
("31 Oct 2026","Halloween (a Saturday)","London-wide","—","Family + young leisure","HIGH","Same-day non-nightclub guide 31 Oct"),
("5 Nov 2026","Bonfire Night","London-wide","—","Family, domestic","HIGH","Displays ranked by the walk home, 3 Nov"),
("7–8 Nov 2026","Fireworks & Drone Festival","Alexandra Palace, N22","Alexandra Palace","Family, domestic","HIGH","Inside anchor 2 Nov"),
("8 Nov 2026","England v Australia","Twickenham","Twickenham rail","International, corporate","PEAK","Away-fan pub guide 7 Nov"),
("13 Nov 2026 – 3 Jan 2027","Christmas at Kew","Kew Gardens, TW9","Kew Gardens","Family, local","HIGH","Anchors SW London family demand for 7 weeks"),
("13–17 Nov 2026","Noah Kahan ×5 nights","The O2, SE10","North Greenwich","Young leisure","PEAK","Anchor 9 Nov + timed walk to the O2, 10 Nov"),
("14 Nov 2026","England v Japan","Twickenham","Twickenham rail","International","HIGH","Inside anchor 9 Nov"),
("19 Nov 2026 – 3 Jan 2027","HYDE PARK WINTER WONDERLAND","Hyde Park, W1/W2","Hyde Park Corner","Family, international","PEAK ★","Anchor 16 Nov + insider guide 17 Nov. 46 days of Zone 1 footfall"),
("21 Nov 2026","England v New Zealand","Twickenham","Twickenham rail","International, corporate","PEAK ★","Timed Waterloo→Twickenham journey, 21 Nov. Strongest ADR day of Nov"),
("27–29 Nov 2026","NATIONS CHAMPIONSHIP FINALS","Twickenham","Twickenham rail","12 nations, international","PEAK ★★","Anchor 23 Nov + availability post 24 Nov + live coverage 28 Nov"),
("26 Nov – 18 Dec 2026","Corporate Christmas party season","City, Canary Wharf, Shoreditch","Various","Corporate","PEAK","Anchors 30 Nov + group-booking post 1 Dec"),
("11 Dec 2026 – 3 Jan 2027","World Darts Championship","Alexandra Palace, N22","Alexandra Palace","Domestic groups","PEAK","Anchor 7 Dec + N London journey post 8 Dec. 180,000 tickets, 24 days"),
("17–21 Dec 2026","London International Horse Show","ExCeL, E16","Custom House","Family, county","HIGH","Pre-Christmas Docklands family demand as corporate fades"),
("31 Dec 2026","New Year's Eve fireworks","South Bank / Embankment","Embankment","International","PEAK ★","Highest single-night rate of the year — enforce 3-night minimums"),
("4–24 Jan 2027","— THE TROUGH —","London-wide","—","—","TROUGH ★","Three genuinely dead weeks. Next quarter's content must sell value, emptiness and closing exhibitions, not events"),
]
r=5
for e in E:
    for i,v in enumerate(e,1):
        c=ws.cell(row=r,column=i,value=v); c.font=Font(name=F,size=9); c.border=BOX
        c.alignment=Alignment(vertical="top",wrap_text=(i in (2,7)))
    d=e[5]
    col={"PEAK ★★":"F5B7B1","PEAK ★":"F8D7DA","PEAK":"FCE4E4","HIGH":"FFF3CD","TROUGH":"D6E4F0","TROUGH ★":"AFC9E0"}.get(d)
    if col: ws.cell(row=r,column=6).fill=PatternFill("solid",fgColor=col)
    ws.cell(row=r,column=6).font=Font(name=F,size=9,bold=True)
    ws.row_dimensions[r].height=34; r+=1
ws.freeze_panes="A5"
for s in wb.worksheets: s.sheet_view.showGridLines=False
wb.save("/home/claude/seeker-estates/outputs/Seeker-Estates-Content-Engine.xlsx")
print("saved, sheets:",wb.sheetnames)
