import sys; sys.path.insert(0,'/home/claude/seeker-estates/build')
from datetime import date, timedelta
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
F="Arial"; NAVY="1F3A5F"
thin=Side(style="thin",color="CCCCCC"); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
wb=load_workbook("/home/claude/seeker-estates/build/_c2.xlsx")
def header(ws,row,cols,widths):
    for i,(c,w) in enumerate(zip(cols,widths),1):
        cell=ws.cell(row=row,column=i,value=c); cell.font=Font(name=F,bold=True,color="FFFFFF",size=10)
        cell.fill=PatternFill("solid",fgColor=NAVY); cell.alignment=Alignment(vertical="center",horizontal="center",wrap_text=True)
        cell.border=BOX; ws.column_dimensions[get_column_letter(i)].width=w
    ws.row_dimensions[row].height=32
def title(ws,t,s,n):
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=n); c=ws.cell(row=1,column=1,value=t)
    c.font=Font(name=F,bold=True,size=15,color=NAVY); ws.row_dimensions[1].height=26
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=n); c=ws.cell(row=2,column=1,value=s)
    c.font=Font(name=F,size=9,italic=True,color="666666"); ws.row_dimensions[2].height=16
def rows(ws,data,start,wrap,h=40):
    r=start
    for row in data:
        for i,v in enumerate(row,1):
            c=ws.cell(row=r,column=i,value=v); c.font=Font(name=F,size=9); c.border=BOX
            c.alignment=Alignment(vertical="top",wrap_text=(i in wrap))
        ws.row_dimensions[r].height=h; r+=1
    return r-1

# ---------- SHEET 3: THE SYSTEM ----------
ws=wb.create_sheet("The System")
cols=["Day","Slot name","Pillar","What it is","Why it exists","Platform","Non-negotiable rule"]
widths=[10,24,12,50,54,14,44]
title(ws,"The weekly rhythm","Five posts a week. The same five slots, every week. Consistency of SLOT is what builds a habit in the audience — the content varies, the shape never does.",len(cols))
header(ws,4,cols,widths)
S=[
("MONDAY","LONDON THIS WEEK","LONDON","The week's events in 30 seconds, with tube stops. Same title card, same music, every single Monday.","This is the show. It's the reason someone follows an accommodation company they have no intention of booking with yet. It also does the heavy lifting on reach, which everything else borrows from.","IG + FB + TikTok","Never skip it. A missed Monday costs more than a bad Monday."),
("TUESDAY","THE PRODUCT","STAY / PRACTICAL","Either a property (tours, comparisons, timed walks) or a genuinely useful London how-to.","The only slot that sells directly. Alternating it with pure utility stops the feed becoming a property listing, which is the failure mode that kills these accounts.","IG + FB + TikTok","Never two property posts in consecutive Tuesdays."),
("WEDNESDAY","LANDLORD DAY","LANDLORD","Regulation, economics, and honest advice for London landlords. No pitch.","The second business objective gets its own protected slot so it never competes with guest content for space. 13 weeks of this makes you the account London landlords check.","IG + FB (skip TikTok)","Education only. The CTA is a save or a DM, never a sales line."),
("THURSDAY","PROOF","PROOF","Reviews (including bad ones), the turnover, the systems, guest stories, what you fixed.","Short-lets have a trust problem — travellers genuinely fear the listing is fake or filthy. Showing the boring reality answers that better than any five-star screenshot.","IG + FB + TikTok","Publish at least one unflattering thing a month. That's the whole mechanism."),
("SATURDAY","DISCOVER","LONDON","Hidden gems, weekend plans, free things, honest rankings, what's open.","Saturday is when people plan. This is the most shareable slot — content that gets sent into a group chat, which is the highest-value share there is.","IG + TikTok","Post by 8am. A Saturday post at 4pm has missed the plan."),
("FRI + SUN","STORIES ONLY","-","No feed post. Stories: polls, behind-the-scenes, resharing guest content, answering DMs publicly.","Two rest days protects quality on the other five and stops the account feeling like a broadcast.","Stories","No feed posts. Protect the rhythm."),
("MONTHLY","THE HUMAN ONE","HUMAN","Founder story, the team, why you do this, what went wrong.","People book from people, and landlords hand over £600k assets to people. Roughly one a month, dropped into any slot.","IG + FB + TikTok","Keep the failures in. The failure is the interesting part."),
]
last=rows(ws,S,5,{4,5,7},76)

r=last+3
ws.cell(row=r,column=1,value="THE SIX PILLARS").font=Font(name=F,bold=True,size=13,color=NAVY)
r+=1
header(ws,r,["Pillar","Share","What it covers","The job it does","Failure mode if you overdo it"],[14,9,50,54,44])
PL=[
("LONDON","43%","Events, neighbourhoods, hidden gems, what's on, free things, honest rankings","Reach. This is what strangers find, follow and share. It is the top of everything.","None, really — this is the safe one to over-index on."),
("LANDLORD","19%","Regulation, economics, contracts, what to check, market data","The second revenue line. Landlord acquisition is worth far more per conversion than a booking.","Guest audience tunes out. Solved by keeping it to Wednesdays."),
("PROOF","18%","Reviews, turnovers, systems, guest stories, fixes, transparency","Trust. Converts the people reach brought you.","Starts to feel defensive. Keep it factual, not pleading."),
("STAY","13%","Property tours, comparisons, timed walks, availability","Sells. The direct-response pillar.","THE killer. An account that is mostly its own properties is an advert, and adverts don't get followed. Hard-capped at ~1 in 7 posts."),
("PRACTICAL","4%","How-tos, guides, problem-solvers, the 6am Heathrow post","Saves and long-tail reach. These posts earn for years.","Can drift off-brand into generic travel content. Keep London-specific."),
("HUMAN","3%","Founder, team, story, mistakes","Makes the brand a person. Disproportionately important for landlord conversion.","Becomes self-indulgent. Once a month is plenty."),
]
r+=1; last2=rows(ws,PL,r,{3,4,5},58)

r=last2+3
ws.cell(row=r,column=1,value="THE APPROVAL LOOP — how a post gets from idea to published").font=Font(name=F,bold=True,size=13,color=NAVY)
r+=1
header(ws,r,["Step","Who","What happens","Time","Rule"],[8,20,70,14,50])
AP=[
("1","Content agent","Drafts the week's five posts from the calendar: hook, script, caption, hashtags, CTA, and the asset list each one needs.","Fri, 30 min","Drafts a full week at once, never one post at a time."),
("2","You (Corentin)","Read the five. Approve, kill or send back. Nothing goes out unapproved.","Fri, 15 min","This is the ONLY human gate. Don't add more."),
("3","Whoever's shooting","Captures the week's assets from the Asset Requests tab.","Sat/Sun, ~3 hrs","Shoot for the week ahead, never the week you're in."),
("4","Content agent","Assembles, schedules to Metricool for Mon-Sat.","Sun, 20 min","IG + FB publish automatically. TikTok lands as a draft — someone taps publish."),
("5","You or ops","Reply to every comment in the first 30 minutes after each post.","5 min x 5","The cheapest reach lever that exists. Non-negotiable."),
("6","Analytics agent","Monday morning: last week's numbers, what worked, what to change.","Mon, auto","Feeds next Friday's drafting. This is the loop closing."),
]
r+=1; rows(ws,AP,r,{3,5},46)
wb.save("/home/claude/seeker-estates/build/_c3.xlsx")
print("system sheet done")
