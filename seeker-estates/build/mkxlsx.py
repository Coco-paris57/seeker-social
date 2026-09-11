import sys; sys.path.insert(0,'/home/claude/seeker-estates/build')
from posts import P
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

F="Arial"
NAVY="1F3A5F"; SAND="F5F1E8"; INK="222222"
PILL={"LONDON":"D6E4F0","STAY":"E4EFE0","LANDLORD":"F7E4D6","PROOF":"EDE0F0","PRACTICAL":"FDF2CC","HUMAN":"F5DCDC"}
thin=Side(style="thin",color="CCCCCC"); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)

wb=Workbook(); 

def header(ws,row,cols,widths):
    for i,(c,w) in enumerate(zip(cols,widths),1):
        cell=ws.cell(row=row,column=i,value=c)
        cell.font=Font(name=F,bold=True,color="FFFFFF",size=10)
        cell.fill=PatternFill("solid",fgColor=NAVY)
        cell.alignment=Alignment(vertical="center",horizontal="center",wrap_text=True)
        cell.border=BOX
        ws.column_dimensions[get_column_letter(i)].width=w
    ws.row_dimensions[row].height=32

def title(ws,text,sub,ncols):
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=ncols)
    c=ws.cell(row=1,column=1,value=text); c.font=Font(name=F,bold=True,size=15,color=NAVY)
    c.alignment=Alignment(vertical="center"); ws.row_dimensions[1].height=26
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=ncols)
    c=ws.cell(row=2,column=1,value=sub); c.font=Font(name=F,size=9,italic=True,color="666666")
    ws.row_dimensions[2].height=16

# ---------- SHEET 1: CONTENT CALENDAR ----------
ws=wb.active; ws.title="Content Calendar"
cols=["Date","Day","Wk","Pillar","Platform","Format","HOOK (first 2 seconds)","Concept / what to film","Why this post / direction","CTA","Assets needed","Event tie","Status","Published?","Reach","Saves"]
widths=[11,6,5,11,10,11,44,52,56,26,34,26,13,11,9,8]
title(ws,"Seeker Estates — 90-Day Content Calendar","10 Sep – 12 Dec 2026 · 68 posts · 5 per week (Mon/Tue/Wed/Thu/Sat) · Fill the yellow columns as you go",len(cols))
header(ws,4,cols,widths)
r=5
for p in P:
    d=date.fromisoformat(p["date"])
    wk=(d-date(2026,9,14)).days//7+1
    vals=[d,d.strftime("%a"),(f"W{wk}" if wk>0 else "W0"),p["pillar"],p["plat"],p["fmt"],p["hook"],
          p["concept"],p["caption"],p["cta"],p["assets"],p["event"],"Not started","","",""]
    for i,v in enumerate(vals,1):
        c=ws.cell(row=r,column=i,value=v)
        c.font=Font(name=F,size=9); c.border=BOX
        c.alignment=Alignment(vertical="top",wrap_text=(i in (7,8,9,10,11,12)))
        if i==1: c.number_format="ddd dd mmm"
        if i==4: c.fill=PatternFill("solid",fgColor=PILL[p["pillar"]]); c.font=Font(name=F,size=9,bold=True)
        if i in (13,14,15,16): c.fill=PatternFill("solid",fgColor="FFFFCC")
    ws.row_dimensions[r].height=74
    r+=1
last=r-1
dv=DataValidation(type="list",formula1='"Not started,Assets needed,Drafting,Awaiting approval,Scheduled,Published"',allow_blank=True)
ws.add_data_validation(dv); dv.add(f"M5:M{last}")
dv2=DataValidation(type="list",formula1='"Yes,No"',allow_blank=True)
ws.add_data_validation(dv2); dv2.add(f"N5:N{last}")
ws.freeze_panes="G5"
ws.auto_filter.ref=f"A4:P{last}"

# summary block
sr=last+2
ws.cell(row=sr,column=1,value="PILLAR MIX").font=Font(name=F,bold=True,size=11,color=NAVY)
for i,(k,lab) in enumerate([("LONDON","Place & events — the reach engine"),("LANDLORD","B2B — second objective"),
        ("PROOF","Trust, reviews, behind-the-scenes"),("STAY","The properties themselves"),
        ("PRACTICAL","Utility / save-worthy"),("HUMAN","People behind the business")],1):
    ws.cell(row=sr+i,column=1,value=k).font=Font(name=F,size=9,bold=True)
    ws.cell(row=sr+i,column=1).fill=PatternFill("solid",fgColor=PILL[k])
    ws.cell(row=sr+i,column=2,value=f'=COUNTIF(D$5:D${last},A{sr+i})').font=Font(name=F,size=9)
    ws.cell(row=sr+i,column=3,value=f'=B{sr+i}/COUNTA(D$5:D${last})').font=Font(name=F,size=9)
    ws.cell(row=sr+i,column=3).number_format="0.0%"
    ws.cell(row=sr+i,column=4,value=lab).font=Font(name=F,size=9,italic=True,color="666666")
ws.cell(row=sr+7,column=1,value="TOTAL").font=Font(name=F,size=9,bold=True)
ws.cell(row=sr+7,column=2,value=f'=COUNTA(D$5:D${last})').font=Font(name=F,size=9,bold=True)
ws.cell(row=sr+9,column=1,value="Published so far").font=Font(name=F,size=9,bold=True)
ws.cell(row=sr+9,column=2,value=f'=COUNTIF(N$5:N${last},"Yes")').font=Font(name=F,size=9)
ws.cell(row=sr+10,column=1,value="Avg reach (published)").font=Font(name=F,size=9,bold=True)
ws.cell(row=sr+10,column=2,value=f'=IFERROR(AVERAGE(O$5:O${last}),0)').font=Font(name=F,size=9)
ws.cell(row=sr+11,column=1,value="Total saves").font=Font(name=F,size=9,bold=True)
ws.cell(row=sr+11,column=2,value=f'=SUM(P$5:P${last})').font=Font(name=F,size=9)
print("calendar rows:",last-4)
wb.save("/home/claude/seeker-estates/build/_c1.xlsx")
